import io

import docx
import openpyxl
from pypdf import PdfWriter

from lad_outlook_mcp.extract import extract_text, html_to_text


def test_docx():
    d = docx.Document()
    d.add_paragraph("Hallo Welt")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "A", "B"
    buf = io.BytesIO()
    d.save(buf)
    text = extract_text("x.docx", None, buf.getvalue())
    assert "Hallo Welt" in text and "A | B" in text


def test_xlsx():
    wb = openpyxl.Workbook()
    wb.active.append(["Artikel", "Preis"])
    wb.active.append(["Krone", 199])
    buf = io.BytesIO()
    wb.save(buf)
    text = extract_text("x.xlsx", None, buf.getvalue())
    assert "Artikel | Preis" in text and "Krone | 199" in text


def test_pdf_blank_page():
    w = PdfWriter()
    w.add_blank_page(100, 100)
    buf = io.BytesIO()
    w.write(buf)
    assert "Seite 1" in extract_text("x.pdf", "application/pdf", buf.getvalue())


def test_text_and_unknown():
    assert extract_text("a.csv", None, "ä;b".encode("cp1252")) == "ä;b"
    assert extract_text("a.zip", "application/zip", b"PK") is None


def test_html_to_text():
    assert html_to_text("<p>Hallo&nbsp;<b>du</b></p><script>x()</script><p>zwei</p>") == "Hallo du\nzwei"
