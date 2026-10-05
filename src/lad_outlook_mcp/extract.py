"""Textextraktion aus Anhängen (PDF, Word, Excel, Text, HTML)."""

from __future__ import annotations

import io
import re
from html import unescape
from pathlib import PurePath

TEXT_EXTENSIONS = {".txt", ".csv", ".tsv", ".json", ".xml", ".md", ".log", ".ics", ".vcf", ".eml"}


def html_to_text(html: str) -> str:
    text = re.sub(r"(?is)<(script|style|head)\b.*?</\1>", " ", html)
    text = re.sub(r"(?i)<br\s*/?>|</p>|</div>|</tr>|</li>|</h\d>", "\n", text)
    text = re.sub(r"(?s)<[^>]+>", " ", text)
    text = unescape(text).replace("\xa0", " ")
    lines = (re.sub(r"[ \t\r\f\v]+", " ", line).strip() for line in text.split("\n"))
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for i, page in enumerate(reader.pages, 1):
        pages.append(f"--- Seite {i} ---\n{(page.extract_text() or '').strip()}")
    return "\n\n".join(pages)


def _docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for t_index, table in enumerate(document.tables, 1):
        parts.append(f"--- Tabelle {t_index} ---")
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _xlsx(data: bytes, max_rows: int = 500) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"--- Blatt: {ws.title} ---")
        for n, row in enumerate(ws.iter_rows(values_only=True)):
            if n >= max_rows:
                parts.append(f"... (gekürzt nach {max_rows} Zeilen)")
                break
            if any(v is not None for v in row):
                parts.append(" | ".join("" if v is None else str(v) for v in row))
    wb.close()
    return "\n".join(parts)


def extract_text(name: str, content_type: str | None, data: bytes) -> str | None:
    """Gibt den Text zurück oder None, wenn das Format nicht unterstützt wird."""
    ext = PurePath(name or "").suffix.lower()
    ctype = (content_type or "").lower()
    if ext == ".pdf" or ctype == "application/pdf":
        return _pdf(data)
    if ext == ".docx" or "wordprocessingml" in ctype:
        return _docx(data)
    if ext in (".xlsx", ".xlsm") or "spreadsheetml" in ctype:
        return _xlsx(data)
    if ext in (".html", ".htm") or ctype == "text/html":
        return html_to_text(_decode(data))
    if ext in TEXT_EXTENSIONS or ctype.startswith("text/"):
        return _decode(data)
    return None
